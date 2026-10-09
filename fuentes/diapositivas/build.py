#!/usr/bin/env python3
"""Build the presentation on top of the university PowerPoint template.

Usage: python fuentes/diapositivas/build.py
Requires: python-pptx. Output: entrega/diapositivas_tanque.pptx
"""
import io
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent.parent
TEMPLATE = ROOT / "fuentes/plantillas/power_point_65.pptx"
CAPTURES = ROOT / "fuentes/capturas"
ASSIGNMENT = ROOT / "docs/Modelo matematico del Tanque.docx"
OUTPUT = ROOT / "entrega/diapositivas_tanque.pptx"

RED = RGBColor(0xB0, 0x35, 0x35)
NAVY = RGBColor(0x1F, 0x38, 0x64)
INK = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
CARD = RGBColor(0xF2, 0xF2, 0xF2)
CARD_LINE = RGBColor(0xD9, 0xD9, 0xD9)
TINT = RGBColor(0xFB, 0xF1, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SERIF = "Noto Serif"
SANS = "Carlito"

LAYOUT_TITLE, LAYOUT_CONTENT, LAYOUT_CLOSING = 0, 1, 2

MARKUP = re.compile(r"(\*\*.+?\*\*|_\{.+?\}|\^\{.+?\})")


def capture(name):
    return str(CAPTURES / name)


def add_runs(paragraph, text, size, color, bold=False, font=SANS, italic=False):
    """Add text with **bold**, _{subscript} and ^{superscript} markup."""
    for part in MARKUP.split(text):
        if not part:
            continue
        run = paragraph.add_run()
        run_bold, baseline = bold, None
        if part.startswith("**"):
            part, run_bold = part[2:-2], True
        elif part.startswith("_{"):
            part, baseline = part[2:-1], "-25000"
        elif part.startswith("^{"):
            part, baseline = part[2:-1], "30000"
        run.text = part
        run.font.size = Pt(size)
        run.font.bold = run_bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = color
        if baseline:
            run._r.get_or_add_rPr().set("baseline", baseline)


def text(slide, x, y, w, h, lines, size=14, color=INK, bold=False, font=SANS, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullets=False, space_after=6, italic=False, spacing=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(space_after)
        if bullets:
            ppr = paragraph._p.get_or_add_pPr()
            ppr.set("marL", "228600")
            ppr.set("indent", "-228600")
            bu = ppr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
            ppr.append(bu)
        add_runs(paragraph, line, size, color, bold, font, italic)
        if spacing:
            for run in paragraph.runs:
                run._r.get_or_add_rPr().set("spc", str(spacing))
    return box


def card(slide, x, y, w, h, title=None, body=None, emphasis=False, body_size=14, bullets=False, title_size=15,
         space_after=6):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.04
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = TINT if emphasis else CARD
    shape.line.color.rgb = RED if emphasis else CARD_LINE
    shape.line.width = Pt(1.5 if emphasis else 0.75)
    top = y + 0.16
    if title:
        text(slide, x + 0.22, top, w - 0.44, 0.34, title, size=title_size, color=RED, bold=True, font=SERIF)
        top += 0.42
    if body:
        text(slide, x + 0.22, top, w - 0.44, h - (top - y) - 0.12, body, size=body_size, bullets=bullets,
             space_after=space_after)
    return shape


def badge(slide, x, y, label, color=NAVY, d=0.42):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_runs(frame.paragraphs[0], str(label), 13, WHITE, bold=True)


def pill(slide, x, y, w, h, label, size=13):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.12
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = NAVY
    shape.line.fill.background()
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = Inches(0.06)
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_runs(frame.paragraphs[0], label, size, WHITE, bold=True)


def picture(slide, source, x, y, w):
    pic = slide.shapes.add_picture(source, Inches(x), Inches(y), width=Inches(w))
    pic.line.color.rgb = CARD_LINE
    pic.line.width = Pt(0.75)
    return pic


def content_slide(prs, kicker, title, number, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    for placeholder in list(slide.placeholders):
        placeholder._element.getparent().remove(placeholder._element)
    text(slide, 0.6, 0.42, 10.5, 0.25, kicker.upper(), size=10, color=NAVY, bold=True, spacing=300)
    text(slide, 0.6, 0.72, 11.6, 0.7, title, size=28, color=RED, bold=True, font=SERIF)
    text(slide, 0.6, 6.78, 1.0, 0.25, str(number), size=10, color=MUTED)
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def numbered_rows(slide, x, y, w, rows, gap=0.78, color=NAVY, label_w=2.6):
    """Rows of badge + bold label + description, like the guide's step lists."""
    for i, (label, description) in enumerate(rows):
        top = y + i * gap
        badge(slide, x, top, i + 1, color)
        text(slide, x + 0.62, top + 0.06, label_w, 0.4, label, size=14, color=RED, bold=True)
        text(slide, x + 0.62 + label_w, top + 0.06, w - 0.62 - label_w, 0.6, description, size=14)


def build():
    prs = Presentation(str(TEMPLATE))
    template_slide_count = len(prs.slides)
    with zipfile.ZipFile(ASSIGNMENT) as z:
        figure1 = z.read("word/media/image1.png")

    # 1. Title
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE])
    for placeholder in list(s.placeholders):
        placeholder._element.getparent().remove(placeholder._element)
    center = dict(align=PP_ALIGN.CENTER)
    text(s, 1.5, 2.72, 10.33, 0.3, "UNIVERSIDAD DE PAMPLONA", size=12, color=NAVY, bold=True, spacing=300, **center)
    text(s, 1.5, 3.12, 10.33, 0.3, "TRABAJO 3", size=11, color=RED, bold=True, spacing=300, **center)
    text(s, 1.5, 3.45, 10.33, 1.2, ["Modelado y simulación de un tanque", "de distribución de agua"],
         size=32, color=RED, bold=True, font=SERIF, space_after=0, **center)
    text(s, 1.5, 4.72, 10.33, 0.3, "Modelado y simulación de sistemas continuos  ·  Simulink y superagente de IA",
         size=13, color=NAVY, **center)
    card(s, 2.4, 5.1, 8.53, 1.2)
    text(s, 2.4, 5.2, 8.53, 0.7, ["Yenderson Josué Rangel Martínez  ·  Kevin Sebastian Medina Nava",
                                   "Cristian Julián Camargo García"],
         size=14, color=NAVY, bold=True, space_after=2, **center)
    text(s, 2.4, 5.92, 8.53, 0.3, "9 de octubre de 2026", size=12, color=MUTED, **center)
    s.notes_slide.notes_text_frame.text = (
        "Presentación del equipo y del tema: el tanque de distribución de agua, resuelto en Simulink y con un "
        "superagente de IA."
    )

    # 2. What the assignment asks
    s = content_slide(prs, "Trabajo 3 · Objetivo", "Qué pide el trabajo", 2,
                      "Dos partes: (a) Simulink, H(s) y h(t) con su comparación; (b) simulador con IA, ingeniería "
                      "inversa e informe. Todo se entrega en un PDF con el enlace del simulador.")
    card(s, 0.6, 1.75, 12.1, 1.1, "Objetivo",
         "Modelar, simular y analizar un tanque de distribución de agua por dos caminos y comparar los resultados.",
         emphasis=True, body_size=15)
    card(s, 0.6, 3.1, 5.9, 2.3, "Parte (a) · Simulink",
         ["Simular la ecuación diferencial", "Obtener H(s)", "Obtener h(t) por transformada inversa",
          "Comparar las tres respuestas"], bullets=True, body_size=15)
    card(s, 6.8, 3.1, 5.9, 2.3, "Parte (b) · Superagente de IA",
         ["Generar el simulador con IA", "Hacer la ingeniería inversa", "Documentar prompts y proceso",
          "Publicar un enlace ejecutable"], bullets=True, body_size=15)
    text(s, 0.6, 5.72, 1.6, 0.45, "Entregables", size=15, color=RED, bold=True, font=SERIF, anchor=MSO_ANCHOR.MIDDLE)
    for i, item in enumerate(["Informe y diapositivas en un PDF", "Enlace del simulador", "Exposición de 10 minutos"]):
        pill(s, 2.3 + i * 3.5, 5.72, 3.3, 0.45, item)

    # 3. Mathematical model
    s = content_slide(prs, "Parte A · Modelo", "Modelo matemático del proceso", 3,
                      "El balance de masa y la ley de Torricelli dan la ecuación (10). Es no lineal por la raíz del "
                      "nivel. K = 4.47 es consistente: 4.47 por raíz de 5 da 10 m³/h.")
    picture(s, io.BytesIO(figure1), 0.6, 1.75, 4.3)
    text(s, 0.6, 5.05, 4.3, 0.3, "Figura del enunciado: tanque con entrada Q_{i} y salida por válvula",
         size=11, color=MUTED, italic=True)
    card(s, 5.3, 1.75, 3.55, 1.25, "Balance de masa", "A · dh/dt = Q_{i} − Q_{o}")
    card(s, 9.15, 1.75, 3.55, 1.25, "Ley de Torricelli", "Q_{o} = K · √h")
    card(s, 5.3, 3.25, 7.4, 1.5, "Ecuación del proceso (10)", emphasis=True)
    text(s, 5.3, 3.85, 7.4, 0.6, "dh/dt = Q_{i} / A − (K / A) · √h", size=24, color=NAVY, bold=True,
         align=PP_ALIGN.CENTER)
    chips = ["A = 10 m²", "K = 4.47 m^{2.5}/h", "H = 10 m", "h_{e} = 5 m", "Q_{e} = 10 m³/h"]
    for i, chip in enumerate(chips):
        pill(s, 5.3 + i * 1.5, 5.05, 1.4, 0.45, chip, size=12)
    text(s, 5.3, 5.68, 7.4, 0.5, "No lineal: el caudal de salida depende de la raíz del nivel.",
         size=14, color=NAVY, italic=True)

    # 4. Decisions on the statement
    s = content_slide(prs, "Parte A · Análisis", "Decisiones sobre el enunciado", 4,
                      "El enunciado deja puntos abiertos. La decisión clave es la constante de tiempo: 1 min no es "
                      "compatible con A, K y el equilibrio; el valor que sale de los datos es 10 horas.")
    decisions = [
        ("Ecuación (5)", "Imprime dV/dt; se usa dh/dt, que lleva a la ecuación (10)."),
        ("Unidades", "Tiempo en horas y K en m^{2.5}/h. La densidad se cancela."),
        ("Condiciones", "h(0) = 5 m, escalón de 10 a 12 m³/h y 50 h de simulación."),
        ("Límites físicos", "0 ≤ h ≤ H y raíz protegida con √max(h, 0)."),
    ]
    for i, (label, description) in enumerate(decisions):
        x, y = 0.6 + (i % 2) * 3.85, 1.75 + (i // 2) * 2.2
        card(s, x, y, 3.65, 2.0)
        badge(s, x + 0.2, y + 0.2, i + 1)
        text(s, x + 0.78, y + 0.26, 2.7, 0.35, label, size=15, color=RED, bold=True, font=SERIF)
        text(s, x + 0.22, y + 0.82, 3.2, 1.1, description, size=14)
    card(s, 8.6, 1.75, 4.1, 4.2, "Constante de tiempo", emphasis=True)
    text(s, 8.82, 2.4, 3.66, 0.75, "10 h", size=44, color=NAVY, bold=True, font=SERIF)
    text(s, 8.82, 3.2, 3.66, 0.35, "y no 1 min, como dice el enunciado", size=14, color=MUTED)
    text(s, 8.82, 3.75, 3.66, 2.1,
         ["τ = 2A√h_{e} / K ≈ 10 h", "Con 1 min haría falta K ≈ 2683 y un caudal de equilibrio de 6000 m³/h.",
          "La ecuación (10) solo depende de A y K."], size=14, space_after=8)

    # 5. Linearization and transfer function
    s = content_slide(prs, "Parte A · H(s)", "Linealización y función de transferencia", 5,
                      "Laplace solo aplica a ecuaciones lineales, así que se linealiza alrededor de 5 m con Taylor. "
                      "Resultado: primer orden con ganancia 1 y constante de tiempo de 10 h.")
    numbered_rows(s, 0.6, 1.8, 7.6, [
        ("Punto de operación", "h_{e} = 5 m, Q_{e} = 10 m³/h"),
        ("Serie de Taylor", "√h ≈ √h_{e} + (h − h_{e}) / (2√h_{e})"),
        ("Variables de desviación", "h′ = h − h_{e},  q′ = Q_{i} − Q_{e}"),
        ("Laplace", "H(s) = (1/A) / (s + K / (2A√h_{e}))"),
        ("Transformada inversa", "Escalón de 2 m³/h en la entrada"),
    ], gap=0.86)
    card(s, 8.6, 1.75, 4.1, 4.4, "Resultado", emphasis=True)
    text(s, 8.82, 2.4, 3.66, 0.3, "Función de transferencia", size=12, color=MUTED)
    text(s, 8.82, 2.72, 3.66, 0.55, "H(s) = 0.1 / (s + 0.1)", size=22, color=NAVY, bold=True)
    text(s, 8.82, 3.6, 3.66, 0.3, "Respuesta en el tiempo", size=12, color=MUTED)
    text(s, 8.82, 3.92, 3.66, 0.55, "h(t) = 5 + 2(1 − e^{−0.1t})", size=22, color=NAVY, bold=True)
    text(s, 8.82, 4.85, 3.66, 1.1, ["Primer orden", "Ganancia de 1 m por cada m³/h", "τ = 10 h"],
         size=14, bullets=True, space_after=4)

    # 6. Simulink
    s = content_slide(prs, "Parte A · Simulink", "Simulación en Simulink", 6,
                      "Un mismo escalón alimenta dos ramas: arriba la ecuación diferencial no lineal, abajo la función "
                      "de transferencia sobre la desviación. Se simuló de 0 a 50 horas.")
    picture(s, capture("diagrama_bloques.jpeg"), 0.6, 1.75, 7.5)
    text(s, 0.6, 5.78, 7.5, 0.3, "Diagrama de bloques: 0 a 50 h, solucionador de paso variable",
         size=11, color=MUTED, italic=True)
    card(s, 8.45, 1.75, 4.25, 2.05, "Rama superior",
         ["Ecuación diferencial no lineal", "Integrador con h(0) = 5 m", "Realimentación K · √h"],
         bullets=True, body_size=14)
    card(s, 8.45, 4.0, 4.25, 2.05, "Rama inferior",
         ["Función de transferencia", "Entra la desviación Q_{i} − 10", "A la salida se suman 5 m"],
         bullets=True, body_size=14)

    # 7. Responses
    s = content_slide(prs, "Parte A · Resultados", "Tres respuestas, dos comportamientos", 7,
                      "Amarillo: ecuación no lineal. Azul: función de transferencia. La gráfica de MATLAB es la "
                      "analítica y coincide con la azul, porque son el mismo modelo lineal.")
    picture(s, capture("scope_no_lineal_y_func_trans.jpeg"), 0.6, 1.75, 6.4)
    text(s, 0.6, 5.25, 6.4, 0.5, "Scope: no lineal (amarillo) y función de transferencia (azul)",
         size=11, color=MUTED, italic=True)
    picture(s, capture("respuesta_analitica_matlab.jpeg"), 7.3, 1.75, 5.4)
    text(s, 7.3, 4.17, 5.4, 0.3, "MATLAB: h(t) analítica, 5 + 2(1 − e^{−0.1t})", size=11, color=MUTED, italic=True)
    card(s, 7.3, 4.65, 5.4, 1.05, None,
         ["**H(s) y h(t) coinciden** entre sí.", "La **no lineal** se separa al subir el nivel."], body_size=14)

    # 8. Comparison
    s = content_slide(prs, "Parte A · Análisis", "Dónde falla el modelo lineal", 8,
                      "Cerca de 5 m las respuestas coinciden. Al alejarse, el lineal se queda corto: 7.00 frente a "
                      "7.21 m. La resistencia de la válvula crece con el nivel.")
    rows = [("Tiempo (h)", "No lineal (m)", "Lineal (m)", "Diferencia (m)"),
            ("0", "5.000", "5.000", "0.000"), ("10", "6.292", "6.264", "0.027"),
            ("20", "6.816", "6.729", "0.087"), ("30", "7.038", "6.900", "0.138"),
            ("50", "7.175", "6.987", "0.189"), ("Final", "7.207", "7.000", "0.207")]
    table = s.shapes.add_table(len(rows), 4, Inches(0.6), Inches(1.75), Inches(6.9), Inches(3.5)).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (CARD if r % 2 == 0 else WHITE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            add_runs(paragraph, value, 14, WHITE if r == 0 else INK, bold=(r == 0 or r == len(rows) - 1))
    text(s, 0.6, 5.5, 6.9, 0.7, "La válvula opone más resistencia al subir el nivel: el modelo lineal subestima el "
         "nivel final y la lentitud del proceso.", size=14, color=NAVY, italic=True)
    stats = [("7.21 m", "Nivel final no lineal"), ("7.00 m", "Nivel final lineal"),
             ("0.21 m", "Diferencia: 2.9 % del nivel final")]
    for i, (value, label) in enumerate(stats):
        y = 1.75 + i * 1.5
        card(s, 7.9, y, 4.8, 1.3, emphasis=(i == 2))
        text(s, 8.15, y + 0.2, 2.2, 0.9, value, size=32, color=RED if i == 2 else NAVY, bold=True, font=SERIF,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, 10.35, y + 0.2, 2.2, 0.9, label, size=14, anchor=MSO_ANCHOR.MIDDLE)

    # 9. AI super agent process
    s = content_slide(prs, "Parte B · Proceso", "Simulación con superagente de IA", 9,
                      "Primero se analizó el enunciado y se tomaron las decisiones; después se escribió un único "
                      "prompt para v0. El código se validó contra la teoría antes de aceptarlo.")
    steps = ["Análisis", "Contraste", "Prompt", "Generación", "Validación", "Despliegue"]
    for i, step in enumerate(steps):
        pill(s, 0.6 + i * 2.05, 1.8, 1.75, 0.55, step)
        if i < len(steps) - 1:
            text(s, 0.6 + i * 2.05 + 1.75, 1.8, 0.3, 0.55, "→", size=18, color=RED, bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    card(s, 0.6, 2.75, 6.0, 3.4, "Herramientas",
         ["**Claude:** análisis del enunciado y plan", "**v0:** genera el programa de simulación",
          "**Agente local:** integra, valida y corrige", "**Vercel:** publica el simulador"], bullets=True, body_size=16, space_after=14)
    card(s, 6.9, 2.75, 5.8, 3.4, "Un solo prompt para v0", emphasis=True)
    text(s, 7.12, 3.4, 5.36, 1.5, ["Física, parámetros, valores por defecto y requisitos de la interfaz.",
                                    "Las decisiones ya estaban tomadas, y por eso el resultado fue correcto."],
         size=14, space_after=8)
    text(s, 7.12, 5.1, 1.3, 0.8, "18", size=36, color=NAVY, bold=True, font=SERIF, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.2, 5.1, 4.2, 0.8, "prompts significativos documentados en el informe", size=14,
         anchor=MSO_ANCHOR.MIDDLE)

    # 10. Reverse engineering
    s = content_slide(prs, "Parte B · Ingeniería inversa", "Ingeniería inversa del simulador", 10,
                      "El cálculo está separado de la interfaz. Usa RK4 o Euler con paso fijo de 3 minutos. Calcula "
                      "tres curvas, y la lineal y la analítica deben coincidir: eso verifica el integrador.")
    card(s, 0.6, 1.75, 3.85, 2.2, "Arquitectura",
         ["Next.js, React y TypeScript", "Sin servidor: todo en el navegador", "Cálculo separado de la interfaz"],
         bullets=True)
    card(s, 4.75, 1.75, 3.85, 2.2, "Método numérico",
         ["RK4 por defecto, o Euler", "Paso fijo de 0.05 h", "Resuelve todo y luego lo reproduce"], bullets=True)
    card(s, 8.9, 1.75, 3.8, 2.2, "Tres curvas",
         ["No lineal: integrada y limitada", "Lineal: integrada", "Analítica: fórmula cerrada"], bullets=True)
    card(s, 0.6, 4.25, 12.1, 1.6, "Verificación contra la teoría", emphasis=True)
    text(s, 0.82, 4.87, 11.66, 0.9,
         "Ecuación (10), polo 0.1, τ = 10 h y equilibrio (Q/K)² coinciden con la teoría. Observaciones: el campo "
         "dischargeCoefficient guarda K, y el punto de operación está fijo en 5 m.", size=15)

    # 11. Simulation program
    s = content_slide(prs, "Parte B · Programa", "Programa de simulación", 11,
                      "Demostración en vivo si hay tiempo: ejecutar con los valores por defecto y mostrar que la curva "
                      "no lineal llega a 7.21 m y la lineal a 7.00 m.")
    card(s, 0.6, 1.75, 6.6, 1.3, "Enlace para ejecutarlo", emphasis=True)
    link = text(s, 0.82, 2.35, 6.16, 0.5, "https://tanque-six.vercel.app/", size=20, color=NAVY, bold=True)
    link.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://tanque-six.vercel.app/"
    link.text_frame.paragraphs[0].runs[0].font.color.rgb = NAVY
    numbered_rows(s, 0.6, 3.4, 6.6, [
        ("Ajustar", "A, K, H, h(0), caudal, escalón y horizonte"),
        ("Elegir", "Método RK4 o Euler"),
        ("Ejecutar", "El tanque se llena de forma animada"),
        ("Comparar", "No lineal, lineal y analítica"),
    ], gap=0.7, color=RED, label_w=1.3)
    card(s, 7.6, 1.75, 5.1, 4.4, "Qué muestra",
         ["Tanque con el nivel animado", "Gráfica con las tres curvas", "Ecuaciones y supuestos del modelo",
          "Insignia de equilibrio: con Q_{i} = 10 el nivel converge a 5 m",
          "Nota sobre el error de linealización"], bullets=True, body_size=15, space_after=14)

    # 12. Conclusions
    s = content_slide(prs, "Cierre", "Conclusiones", 12,
                      "Cerrar con las cinco ideas: modelo correcto, H(s) de primer orden, límite del modelo lineal, "
                      "valor del análisis previo para el superagente y verificación por ingeniería inversa.")
    conclusions = [
        "El modelo es correcto y consistente: K = 4.47 da 10 m³/h a 5 m de nivel.",
        "H(s) = 0.1 / (s + 0.1): primer orden con τ = 10 h, y no 1 min.",
        "El modelo lineal sirve cerca de 5 m; lejos se queda corto: 7.00 frente a 7.21 m.",
        "El superagente acertó con un solo prompt porque las decisiones ya estaban tomadas.",
        "La ingeniería inversa confirmó que el código implementa la teoría.",
    ]
    for i, conclusion in enumerate(conclusions):
        y = 1.8 + i * 0.86
        badge(s, 0.6, y, i + 1)
        text(s, 1.25, y, 11.2, 0.5, conclusion, size=16, anchor=MSO_ANCHOR.MIDDLE)

    # 13. Closing
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CLOSING])
    text(s, 1.5, 4.55, 10.33, 0.8, "Gracias", size=40, color=RED, bold=True, font=SERIF, align=PP_ALIGN.CENTER)
    text(s, 1.5, 5.4, 10.33, 0.4, "¿Preguntas?", size=18, color=NAVY, align=PP_ALIGN.CENTER)
    closing_link = text(s, 1.5, 5.9, 10.33, 0.4, "https://tanque-six.vercel.app/", size=16, color=NAVY, bold=True,
                        align=PP_ALIGN.CENTER)
    closing_link.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://tanque-six.vercel.app/"
    s.notes_slide.notes_text_frame.text = "Cierre y preguntas. El enlace del simulador queda visible."

    # Drop the template's sample slides.
    id_list = prs.slides._sldIdLst
    for slide_id in list(id_list)[:template_slide_count]:
        prs.part.drop_rel(slide_id.rId)
        id_list.remove(slide_id)

    prs.save(str(OUTPUT))
    print(f"written {OUTPUT.relative_to(ROOT)} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
